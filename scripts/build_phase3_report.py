from __future__ import annotations

import json
import re
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "CampusEnergyAI_Report.docx"
METRICS = json.loads((ROOT / "artifacts" / "metrics.json").read_text(encoding="utf-8"))
SPLITS = json.loads((ROOT / "artifacts" / "split_metadata.json").read_text(encoding="utf-8"))
FORECAST = pd.read_csv(ROOT / "outputs" / "future_24h_forecast.csv")
SUCCESS = pd.read_csv(ROOT / "outputs" / "success_cases.csv")
FAILURE = pd.read_csv(ROOT / "outputs" / "failure_cases.csv")
IMPORTANCE = pd.read_csv(ROOT / "artifacts" / "feature_importance.csv")
RAW = pd.read_csv(ROOT / "data" / "raw" / "campus_energy_hourly.csv")

SELECTED_MODEL = METRICS["selection"]["selected_model"]
CANDIDATES = METRICS["selection"]["validation_candidates"]
FINAL = METRICS["final_test"]
HIGH_USE = METRICS["high_use"]
FUTURE = METRICS["future_forecast"]
SPLIT = METRICS["split"]

NAVY = "16324F"
TEAL = "159D9A"
PALE = "EAF4F4"
MUTED = "52616B"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = "D8E0E5") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), color)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, end])


def add_title_page(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Inches(1.2)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("CampusEnergy AI")
    r.bold = True
    r.font.size = Pt(34)
    r.font.color.rgb = RGBColor.from_string(NAVY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("24-Hour Campus Energy Forecasting & Decision Support")
    r.font.size = Pt(18)
    r.font.color.rgb = RGBColor.from_string(TEAL)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Inches(0.55)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("HACKATHON PROJECT · TECHNICAL REPORT")
    r.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor.from_string(MUTED)

    image = ROOT / "outputs" / "phase2" / "architecture.png"
    doc.add_picture(str(image), width=Inches(6.6))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    table = doc.add_table(rows=3, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    labels = ["Participant", "Team", "Institution"]
    values = ["[Participant Name]", "[Team Name]", "[Institution]"]
    for i, (label, value) in enumerate(zip(labels, values)):
        table.cell(i, 0).width = Inches(1.5)
        table.cell(i, 1).width = Inches(4.3)
        table.cell(i, 0).text = label
        table.cell(i, 1).text = value
        table.cell(i, 0).paragraphs[0].runs[0].bold = True
        for cell in table.rows[i].cells:
            set_cell_border(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Inches(0.3)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Synthetic / simulated prototype campus energy data")
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor.from_string(MUTED)


def add_heading(doc: Document, title: str, level: int = 1, page_break: bool = False) -> None:
    if page_break:
        doc.add_page_break()
    doc.add_heading(title, level=level)


def add_body(doc: Document, text: str, bold_prefix: str | None = None) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        r.bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)


def add_numbered_reference(doc: Document, text: str) -> None:
    """Render the small Markdown subset used by references.md without raw markers."""
    p = doc.add_paragraph(style="List Number")
    for part in re.split(r"(\*[^*]+\*)", text):
        if not part:
            continue
        run = p.add_run(part[1:-1] if part.startswith("*") and part.endswith("*") else part)
        run.italic = part.startswith("*") and part.endswith("*")


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths=None) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for j, h in enumerate(headers):
        cell = table.rows[0].cells[j]
        cell.text = h
        set_cell_shading(cell, NAVY)
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor.from_string(WHITE)
        set_cell_border(cell, WHITE)
    for i, row in enumerate(rows):
        cells = table.add_row().cells
        for j, value in enumerate(row):
            cells[j].text = str(value)
            if i % 2 == 1:
                set_cell_shading(cells[j], PALE)
            set_cell_border(cells[j])
            cells[j].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for row in table.rows:
        row.allow_break_across_pages = False
    doc.add_paragraph()


def add_figure(doc: Document, relative: str, caption: str, width: float = 6.55) -> None:
    path = ROOT / relative
    doc.add_picture(str(path), width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    r = p.add_run(caption)
    r.italic = True
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor.from_string(MUTED)


def fmt_ts(value: str) -> str:
    return pd.to_datetime(value).strftime("%d %b %Y, %H:%M")


def build() -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    section.header_distance = Inches(0.32)
    section.footer_distance = Inches(0.32)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].paragraph_format.line_spacing = 1.08
    for name, size, color in [("Title", 34, NAVY), ("Heading 1", 20, NAVY), ("Heading 2", 14, TEAL), ("Heading 3", 11.5, NAVY)]:
        styles[name].font.name = "Aptos Display" if name != "Normal" else "Aptos"
        styles[name].font.size = Pt(size)
        styles[name].font.color.rgb = RGBColor.from_string(color)
        styles[name].font.bold = True
    styles["Heading 1"].paragraph_format.space_before = Pt(0)
    styles["Heading 1"].paragraph_format.space_after = Pt(9)
    styles["Heading 2"].paragraph_format.space_before = Pt(10)
    styles["Heading 2"].paragraph_format.space_after = Pt(5)

    header = section.header.paragraphs[0]
    header.text = "CAMPUSENERGY AI  ·  TECHNICAL REPORT"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor.from_string(MUTED)
    add_page_number(section.footer.paragraphs[0])

    add_title_page(doc)

    add_heading(doc, "Executive Summary", page_break=True)
    add_body(doc, "CampusEnergy AI is a decision-support prototype for forecasting hourly campus-building electricity demand one day ahead. It combines a reproducible synthetic dataset, leakage-safe time-series features, chronological model selection, an explicit previous-day benchmark, peak detection, interpretable recommendations, scenario analysis, and a six-page Streamlit dashboard. The intended user is a facility or energy manager who needs advance notice of likely peaks rather than only retrospective reporting.")
    add_body(doc, f"On the untouched {FINAL['test_rows']}-hour test period, the selected RandomForestRegressor achieved MAE {FINAL['ml_model']['mae']:.4f} kWh and RMSE {FINAL['ml_model']['rmse']:.4f} kWh. The exact previous-day same-hour baseline achieved MAE {FINAL['baseline']['mae']:.4f} kWh and RMSE {FINAL['baseline']['rmse']:.4f} kWh, corresponding to a {FINAL['mae_improvement_percent']:.2f}% MAE reduction on this synthetic prototype dataset. Random Forest and HistGradientBoosting were effectively tied on validation; the former was selected only because its canonical validation MAE was marginally lower.")
    add_body(doc, "The next-24-hour forecast covers 30 June 2026 and identifies two HIGH USE hours. Recommendations remain transparent checks for a human manager. The prototype does not use live campus data and does not directly control equipment.")

    add_heading(doc, "Contents", page_break=True)
    sections = [
        "1. Problem Definition", "2. Intended Users", "3. Proposed Solution", "4. Innovation and Key Features",
        "5. System Architecture", "6. Dataset", "7. Target and Features", "8. Preprocessing and Feature Engineering",
        "9. Time-Series Data Split", "10. Data Leakage Prevention", "11. Required Baseline", "12. Machine-Learning Methodology",
        "13. Evaluation Metrics", "14. Final Test Results", "15. Why ML Beat the Baseline", "16. Next-24-Hour Forecast",
        "17. High-Use Detection", "18. Recommendation System", "19. Explainability", "20. What-if Scenario Analysis",
        "21. Working Prototype", "22. Testing and Verification", "23. Representative Successful Cases", "24. Failure / Error Analysis",
        "25. Responsible AI", "26. Limitations", "27. SDG Alignment", "28. Implemented vs Future Work", "29. Conclusion", "References"
    ]
    add_bullets(doc, sections)

    add_heading(doc, "1. Problem Definition", page_break=True)
    add_body(doc, "Campus buildings exhibit time-dependent electricity demand shaped by schedules, weather, recurring weekly patterns, and occasional events. Conventional reports often reveal a peak after it has occurred. That supports accounting, but gives a facility manager little time to inspect HVAC schedules, flexible loads, lighting, or event-related demand before the peak.")
    add_body(doc, "The design question is therefore practical: can a transparent prototype forecast the next 24 hourly values, mark unusually high demand, show why the hour may be high, and recommend specific checks while leaving every operational decision with a human?")

    add_heading(doc, "2. Intended Users")
    add_body(doc, "The primary user is a campus or building facility-energy manager. Secondary users include the facilities team, sustainability team, campus administration, and maintenance staff. Each user needs concise evidence rather than raw model internals: forecast timing, likely peak magnitude, benchmark performance, uncertainty cues, and a defensible next action.")

    add_heading(doc, "3. Proposed Solution")
    add_body(doc, "CampusEnergy AI validates hourly input, creates time and historical-lag features, compares two regression candidates against a required previous-day baseline, and preserves an untouched chronological test period. The selected frozen model produces 24 future predictions. A training-derived high-use threshold marks candidate peaks; deterministic rules translate forecast context into manager checks; the dashboard exposes performance, peaks, scenarios, errors, limitations, and human oversight.")

    add_heading(doc, "4. Innovation and Key Features")
    add_bullets(doc, [
        "An end-to-end forecast-to-decision workflow rather than a standalone ML notebook.",
        "A mandatory, timestamp-aligned previous-day benchmark beside ML results.",
        "Leakage-safe chronological evaluation and an explicitly untouched final test.",
        "Training-only high-use detection, readable indicators, and deterministic recommendations.",
        "Non-causal what-if analysis plus honest successful and difficult cases.",
        "Visible synthetic-data disclosure, Responsible AI boundaries, and human approval before action."
    ])

    add_heading(doc, "5. System Architecture", page_break=True)
    add_figure(doc, "outputs/phase2/architecture.png", "Figure 1. Verified end-to-end architecture. Recommendations stop at human review.", 6.65)
    add_body(doc, "Inputs pass through schema and range validation before feature engineering. Ordered feature rows are divided into training, validation, and test intervals. The two ML candidates are fitted on training data and compared on validation MAE; the test period is used only after selection. Forecast construction uses observed historical lags plus explicitly assumed temperature and occupancy. High-use detection, explanation, and rule-based recommendations feed the facility-manager dashboard.")

    add_heading(doc, "6. Dataset", page_break=True)
    add_body(doc, f"The dataset contains {len(RAW):,} deterministic hourly observations from {fmt_ts(RAW['timestamp'].iloc[0])} to {fmt_ts(RAW['timestamp'].iloc[-1])}. It is synthetic / simulated prototype campus energy data, not real VIT or other campus-meter data. The generator uses fixed seed {METRICS['random_seed']} for reproducibility.")
    add_body(doc, "Consumption combines a base load, weekday and occupancy load, temperature-sensitive cooling, daytime shape, a slow efficiency drift, and random noise. Temperature contains seasonal, diurnal, and autoregressive weather variation. Occupancy follows a simplified weekday and Saturday schedule. Seven labeled event windows add unusual demand; those labels support retrospective error analysis but never enter model features.")
    add_body(doc, "Synthetic data is appropriate here because the challenge permits sample or simulated time-series data and because it enables a complete, privacy-preserving methodology without access to private campus meters. It does not establish real-world accuracy; deployment would require local data, monitoring, and recalibration.")
    add_table(doc, ["Field", "Meaning", "Availability"], [
        ["timestamp", "Hourly observation time", "Known"],
        ["consumption_kwh", "Hourly electricity target", "Observed historically; unknown for future"],
        ["temperature_c", "Simulated outdoor temperature", "Observed historically; assumed for forecast"],
        ["occupied", "Aggregate schedule indicator", "Observed/defined schedule"],
        ["event_label", "Retrospective simulated event annotation", "Excluded from model inputs"]
    ])

    add_heading(doc, "7. Target and Features", page_break=True)
    add_body(doc, "The regression target is consumption_kwh. Model inputs are hour, day_of_week, is_weekend, temperature_c, occupied, lag_24, lag_168, and cyclical sine/cosine encodings for hour and day of week. Hour and calendar fields are known at prediction time. Temperature and occupancy are explicit future assumptions. The lag features refer only to already observed consumption.")
    add_table(doc, ["Feature", "Purpose"], [
        ["hour / cyclical hour", "Daily operating pattern without a false discontinuity at midnight"],
        ["day_of_week / cyclical day", "Weekly schedule pattern"],
        ["is_weekend", "Reduced or different weekend operations"],
        ["temperature_c", "Weather-sensitive cooling demand"],
        ["occupied", "Scheduled occupancy-related load"],
        ["lag_24", "Actual consumption exactly 24 hours earlier"],
        ["lag_168", "Actual consumption exactly one week earlier"]
    ])

    add_heading(doc, "8. Preprocessing and Feature Engineering", page_break=True)
    add_body(doc, "Rows are validated, sorted by timestamp, and required to be unique and hourly. Calendar and cyclical fields are derived from the timestamp. Historical demand enters only through positive shifts: lag_24 equals consumption at t−24 hours and lag_168 equals consumption at t−168 hours. The first 168 hours cannot have a complete weekly lag and are removed after feature creation.")
    add_body(doc, f"This changes {len(RAW):,} raw rows into {SPLIT['rows_after_lags']:,} leakage-safe feature rows. No target value at the prediction timestamp is included in X, and no scaler or encoder is fitted on the complete dataset.")

    add_heading(doc, "9. Time-Series Data Split")
    add_figure(doc, "outputs/charts/chronological_split.png", "Figure 2. Ordered training, validation, and final-test periods.", 6.45)
    add_table(doc, ["Partition", "Range", "Rows", "Purpose"], [
        ["Training", f"{fmt_ts(SPLITS['train']['start'])} – {fmt_ts(SPLITS['train']['end'])}", f"{SPLITS['train']['rows']:,}", "Fit candidate models and threshold"],
        ["Validation", f"{fmt_ts(SPLITS['validation']['start'])} – {fmt_ts(SPLITS['validation']['end'])}", f"{SPLITS['validation']['rows']:,}", "Choose model by MAE"],
        ["Final test", f"{fmt_ts(SPLITS['test']['start'])} – {fmt_ts(SPLITS['test']['end'])}", f"{SPLITS['test']['rows']:,}", "One final evaluation"]
    ])
    add_body(doc, "No random shuffling was used because a forecast model must learn only from the past when predicting later observations. The maximum timestamp of each earlier partition is strictly before the minimum timestamp of the next.")

    add_heading(doc, "10. Data Leakage Prevention", page_break=True)
    add_bullets(doc, [
        "Only positive historical shifts create lag_24 and lag_168.",
        "No negative shift, centered rolling window, backward fill, current target, or future event label enters X.",
        "Chronological iloc slices preserve order; there is no random train_test_split.",
        "Candidate choice and permutation importance use validation data, not the final test.",
        "The test set is evaluated only after the validation winner is fixed.",
        "Future lags use observed history only; no future actual consumption exists in the forecast builder."
    ])
    add_body(doc, "A dedicated integrity audit reconstructed features from raw data, reconciled all timestamps and predictions, and reported: NO LEAKAGE FOUND.")

    add_heading(doc, "11. Required Baseline")
    add_body(doc, "The benchmark predicts each test timestamp with the actual consumption from exactly 24 hours earlier: baseline(t) = consumption(t−24 h). It is intuitive, inexpensive, and demanding enough to detect whether ML adds value beyond daily repetition. Baseline and ML metrics use the identical 624 timestamps.")

    add_heading(doc, "12. Machine-Learning Methodology")
    add_body(doc, "Two tree-based regressors were chosen for nonlinear relationships and mixed calendar, weather, occupancy, and lag features. HistGradientBoostingRegressor builds trees sequentially to correct residual errors. RandomForestRegressor averages many bootstrapped decision trees to reduce variance. Both work well on a compact tabular dataset without the data volume or sequence complexity required by deep learning.")
    candidates = CANDIDATES
    add_table(doc, ["Candidate", "Validation MAE (kWh)", "Validation RMSE (kWh)"], [
        ["HistGradientBoosting", f"{candidates['HistGradientBoostingRegressor']['mae']:.4f}", f"{candidates['HistGradientBoostingRegressor']['rmse']:.4f}"],
        ["Random Forest", f"{candidates['RandomForestRegressor']['mae']:.4f}", f"{candidates['RandomForestRegressor']['rmse']:.4f}"]
    ])
    add_body(doc, "The two candidates performed effectively similarly on validation. Random Forest was selected in the canonical raw-data pipeline because it achieved the marginally lower validation MAE. The final test result was not used to revise that choice.")

    add_heading(doc, "13. Evaluation Metrics")
    add_body(doc, "Mean Absolute Error (MAE) is the average absolute distance between forecast and actual consumption, expressed in kWh and easy to interpret. Root Mean Squared Error (RMSE) squares errors before averaging, so a few large misses have greater influence. Reporting both shows typical accuracy and sensitivity to difficult periods.")

    add_heading(doc, "14. Final Test Results", page_break=True)
    add_table(doc, ["Method", "MAE (kWh)", "RMSE (kWh)", "Rows"], [
        ["Previous-day baseline", f"{FINAL['baseline']['mae']:.4f}", f"{FINAL['baseline']['rmse']:.4f}", f"{FINAL['test_rows']}"],
        ["Selected Random Forest", f"{FINAL['ml_model']['mae']:.4f}", f"{FINAL['ml_model']['rmse']:.4f}", f"{FINAL['test_rows']}"]
    ])
    add_body(doc, f"The selected model reduced MAE by {FINAL['mae_improvement_percent']:.2f}% relative to the baseline. This result is specific to the synthetic prototype dataset and should not be assumed to generalize unchanged to a real campus.")
    add_figure(doc, "outputs/charts/baseline_vs_ml.png", "Figure 3. Final-test comparison using identical timestamps.", 6.45)

    add_heading(doc, "15. Why ML Beat the Baseline", page_break=True)
    add_body(doc, "The baseline assumes that the same hour yesterday is sufficient. It is vulnerable when weekday schedules, occupancy, temperature, or event-related demand differ across consecutive days. The ML model can combine occupancy, weather, calendar position, a 24-hour lag, and a 168-hour weekly lag. This additional context is technically plausible as the source of the measured improvement. The result does not imply that ML must win on every hour or on every real building.")

    add_heading(doc, "16. Next-24-Hour Forecast")
    add_body(doc, f"The frozen forecast contains exactly {len(FORECAST)} consecutive hours from {fmt_ts(FORECAST['timestamp'].iloc[0])} through {fmt_ts(FORECAST['timestamp'].iloc[-1])}. Temperature is an explicit planning assumption based on recent same-hour climatology plus a disclosed warm-day offset. Occupancy follows the documented weekday/Saturday schedule. Both lag features reference observed historical consumption only; no future actual consumption is used.")
    add_figure(doc, "outputs/charts/future_24h_forecast.png", "Figure 4. Canonical next-24-hour forecast and high-use threshold.", 6.45)

    add_heading(doc, "17. High-Use Detection", page_break=True)
    high = FORECAST[FORECAST["high_use"].astype(bool)]
    add_body(doc, f"HIGH USE is defined as a forecast at or above the 90th percentile of training consumption: {HIGH_USE['threshold_kwh']:.3f} kWh. The threshold is not fitted on validation, test, or forecast values. The current forecast contains {len(high)} flagged hours.")
    add_table(doc, ["Forecast time", "Forecast (kWh)", "Status"], [[pd.to_datetime(r.timestamp).strftime("%d %b %Y %H:%M"), f"{r.predicted_kwh:.3f}", "HIGH USE"] for r in high.itertuples()])

    add_heading(doc, "18. Recommendation System")
    add_body(doc, "Recommendations are transparent, rule-based manager checks generated from forecast magnitude, hour, temperature, occupancy, and high-use status. They cover HVAC set-points and schedules, after-hours lighting or plug loads, unusual event/equipment verification, and moving flexible loads when appropriate. The text does not issue control commands.")
    add_body(doc, "The operating boundary is explicit: AI forecast → recommendation → facility-manager verification → human decision. CampusEnergy AI is decision support only and does not automatically switch equipment.")

    add_heading(doc, "19. Explainability")
    add_body(doc, "Validation-set permutation importance measures how much validation MAE increases when one feature is shuffled. In the canonical artifact, occupied and temperature are the strongest measured drivers, with both lag features also contributing. These values describe predictive dependence in this model and dataset; they do not prove that changing a feature will cause the same energy change.")
    add_figure(doc, "outputs/charts/feature_importance.png", "Figure 5. Validation-set permutation importance; larger MAE increase means greater predictive reliance.", 5.1)

    add_heading(doc, "20. What-if Scenario Analysis", page_break=True)
    add_body(doc, "The dashboard lets a user adjust assumed temperature and occupancy for one forecast hour and recompute the model prediction while holding the remaining model inputs fixed. The comparison supports exploration of model sensitivity and operational discussion.")
    add_body(doc, "Scenario analysis — not a guaranteed or causal outcome. The model learned associations from simulated data and does not contain an equipment-level causal model.")
    add_figure(doc, "outputs/phase2/screenshots/what_if.png", "Figure 6. Working what-if interface with the non-causal warning visible.", 6.35)

    add_heading(doc, "21. Working Prototype", page_break=True)
    add_body(doc, "The Streamlit application has six routed pages: Overview, 24-Hour Forecast, Model Performance, Peak Insights, Error Analysis, and Responsible AI / About. What-if scenario controls are part of Peak Insights. Together they expose the full verified workflow without requiring the user to inspect code.")
    add_figure(doc, "outputs/phase2/screenshots/overview.png", "Figure 7. Overview page: forecast summary, performance, disclosure, and boundaries.", 6.45)
    doc.add_page_break()
    add_figure(doc, "outputs/phase2/screenshots/forecast_explorer.png", "Figure 8. Forecast explorer with hourly detail and HIGH USE visibility.", 6.45)
    add_figure(doc, "outputs/phase2/screenshots/performance.png", "Figure 9. Model Performance page with chronological and baseline evidence.", 6.45)

    add_heading(doc, "22. Testing and Verification", page_break=True)
    add_body(doc, "The Phase 2 acceptance run completed 44 pytest tests with zero failures. Coverage includes data validation, chronological splitting, lag identity, baseline alignment, metric recomputation, threshold provenance, future-lag provenance, model contract, deterministic recommendations, scenario behavior, missing-artifact handling, and Streamlit page smoke tests. A real Streamlit server returned HTTP 200, and browser inspection covered navigation, scenario interaction, chart tooltip, error tables, disclosures, and narrow layout behavior.")
    add_body(doc, "A SHA-256 manifest confirms that all 20 frozen Phase 1 data, model, prediction, metric, and chart artifacts remained byte-identical through Phase 2. The report and presentation read those artifacts; they do not retrain the model.")

    add_heading(doc, "23. Representative Successful Cases")
    rows = []
    for r in SUCCESS.head(5).itertuples():
        rows.append([fmt_ts(r.timestamp), f"{r.actual_kwh:.3f}", f"{r.predicted_kwh:.3f}", f"{r.absolute_error:.3f}"])
    add_table(doc, ["Timestamp", "Actual", "Prediction", "Absolute error"], rows)
    add_body(doc, "These rows demonstrate that the model can closely reproduce ordinary operating periods. They are reported as representative evidence, not as a replacement for aggregate metrics.")

    add_heading(doc, "24. Failure / Error Analysis", page_break=True)
    rows = []
    for r in FAILURE.head(6).itertuples():
        label = getattr(r, "event_type", "")
        rows.append([fmt_ts(r.timestamp), f"{r.actual_kwh:.2f}", f"{r.predicted_kwh:.2f}", f"{r.absolute_error:.2f}", str(label).replace("_", " ")])
    add_table(doc, ["Timestamp", "Actual", "Prediction", "Abs. error", "Retrospective context"], rows)
    add_body(doc, "The largest misses occur during simulated lab-load and special-event periods that are not represented in the model inputs. Event labels are shown only after evaluation and were deliberately excluded from training and forecasting. This failure mode is scientifically useful: unusual activity and omitted variables can dominate otherwise strong average performance. A real system should ingest approved schedule information, monitor residuals, and escalate rather than hide large errors.")
    add_figure(doc, "outputs/phase2/screenshots/errors.png", "Figure 10. Error-analysis page with genuine canonical failure cases.", 6.35)

    add_heading(doc, "25. Responsible AI", page_break=True)
    for subtitle, text in [
        ("Privacy", "The prototype uses only aggregate building occupancy, not identities, biometrics, or individual movement. A production deployment should minimize collection and define retention and access policies."),
        ("Bias", "Synthetic schedules may underrepresent holidays, exams, special events, unusual weather, and changed operating patterns. Error monitoring should be segmented by these contexts before operational reliance."),
        ("Security", "Production controls should include authenticated users, least-privilege authorization, encryption, audit logs, secure secrets, and restricted read/write access to meters or a BMS."),
        ("Accessibility", "The dashboard uses readable labels, units, text warnings, and status text rather than color alone. The desktop-first prototype still needs formal screen-reader, keyboard, contrast, and mobile testing."),
        ("Human Oversight", "AI forecast → recommendation → facility-manager verification → human decision. The system never directly controls equipment and cannot replace safety, comfort, or operations procedures.")
    ]:
        add_heading(doc, subtitle, level=2)
        add_body(doc, text)

    add_heading(doc, "26. Limitations")
    add_bullets(doc, [
        "Synthetic, single-building data cannot establish real-campus generalization.",
        "Future temperature and occupancy are assumptions, not live forecasts or sensor feeds.",
        "Simplified occupancy and missing event/equipment context create large errors.",
        "The model learns historical associations, not causal equipment behavior.",
        "There is no live meter or BMS integration and no automatic control.",
        "The horizon is fixed at 24 hours; drift and multi-building scaling are not implemented.",
        "Random Forest and HistGradientBoosting were effectively tied on validation.",
        "The desktop-first UI has limited mobile and formal accessibility validation."
    ])

    add_heading(doc, "27. SDG Alignment", page_break=True)
    add_table(doc, ["Goal", "Project relevance", "Claim boundary"], [
        ["SDG 7 · Affordable and Clean Energy", "Advance demand visibility can support more efficient energy operations.", "No measured energy saving is claimed."],
        ["SDG 11 · Sustainable Cities and Communities", "Data-informed building operations can support more resilient campuses.", "Prototype scope is one simulated building."],
        ["SDG 12 · Responsible Consumption and Production", "Peak and scenario information can support responsible consumption decisions.", "Recommendations require human verification."],
        ["SDG 13 · Climate Action", "Energy-demand awareness may support future emissions-reduction programs.", "No quantified carbon reduction is claimed."]
    ])

    add_heading(doc, "28. Implemented vs Future Work")
    add_table(doc, ["Implemented and verified", "Future production roadmap"], [
        ["Deterministic synthetic hourly data", "Approved real smart-meter ingestion"],
        ["Chronological split and leakage audit", "Live weather and aggregate occupancy feeds"],
        ["Baseline plus two ML candidates", "Multiple-building model registry"],
        ["Frozen 24-hour forecast and peak flags", "Equipment-level anomaly detection"],
        ["Explainability and non-causal scenarios", "Longer horizons and uncertainty intervals"],
        ["Rule-based recommendations and human approval", "Human-approved BMS integration"],
        ["Six-page dashboard and regression tests", "Monitored periodic retraining and drift alerts"]
    ])

    add_heading(doc, "29. Conclusion")
    add_body(doc, "CampusEnergy AI demonstrates a complete, reproducible path from hourly data to a human-reviewed energy decision. Its strongest contribution is the integration of honest benchmarking, leakage-safe validation, peak awareness, explanations, recommendations, failure analysis, and governance in a working dashboard. The measured prototype result is strong, but the report preserves its boundaries: simulated data, assumed future inputs, effectively tied candidates, and no autonomous control. The next meaningful step is a monitored pilot using approved real-building data.")

    add_heading(doc, "References", page_break=True)
    refs = (ROOT / "docs" / "references.md").read_text(encoding="utf-8").splitlines()
    for line in refs:
        if line[:1].isdigit() and ". " in line:
            add_numbered_reference(doc, line.split(". ", 1)[1])

    core_props = doc.core_properties
    core_props.title = "CampusEnergy AI — Technical Report"
    core_props.subject = "24-Hour Campus Energy Forecasting & Decision Support"
    core_props.author = "CampusEnergy AI project team"
    core_props.comments = "Generated from frozen canonical Phase 1 artifacts; no model retraining."

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
