from __future__ import annotations

import json
import re
import zipfile
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / "docs" / "CampusEnergyAI_Report.docx"
PPT = ROOT / "docs" / "CampusEnergyAI_Presentation.pptx"
METRICS = json.loads((ROOT / "artifacts" / "metrics.json").read_text(encoding="utf-8"))


def require(condition: bool, message: str) -> None:
    if not condition:
        raise AssertionError(message)


def xml_text(path: Path, pattern: str) -> str:
    with zipfile.ZipFile(path) as archive:
        parts = []
        for name in archive.namelist():
            if re.fullmatch(pattern, name):
                parts.append(archive.read(name).decode("utf-8", errors="ignore"))
        raw = " ".join(parts)
        return " ".join(re.findall(r">([^<>]+)<", raw))


def validate_report() -> dict:
    require(REPORT.exists() and REPORT.stat().st_size > 100_000, "Report missing or unexpectedly small")
    doc = Document(REPORT)
    text = "\n".join(p.text for p in doc.paragraphs)
    table_text = "\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
    all_text = text + "\n" + table_text
    required = [
        "Executive Summary", "1. Problem Definition", "5. System Architecture", "6. Dataset",
        "9. Time-Series Data Split", "10. Data Leakage Prevention", "11. Required Baseline",
        "12. Machine-Learning Methodology", "14. Final Test Results", "16. Next-24-Hour Forecast",
        "17. High-Use Detection", "20. What-if Scenario Analysis", "21. Working Prototype",
        "22. Testing and Verification", "23. Representative Successful Cases", "24. Failure / Error Analysis",
        "25. Responsible AI", "Privacy", "Bias", "Security", "Accessibility", "Human Oversight",
        "27. SDG Alignment", "28. Implemented vs Future Work", "29. Conclusion", "References",
        "Synthetic / simulated prototype campus energy data", "NO LEAKAGE FOUND", "effectively tied",
        "9.2585", "16.7316", "2.9061", "5.3521", "68.61", "84.882", "4,320", "4,152", "624",
    ]
    for value in required:
        require(value.lower() in all_text.lower(), f"Report missing: {value}")
    lowered = all_text.lower()
    require("todo" not in lowered and "lorem ipsum" not in lowered, "Report contains unfinished text")
    require("*machine learning*" not in lowered and "*scikit-learn documentation*" not in lowered,
            "Report contains raw Markdown formatting in references")
    bracketed = set(re.findall(r"\[[^\]]+\]", all_text))
    require(bracketed <= {"[Participant Name]", "[Team Name]", "[Institution]"}, f"Unexpected placeholders: {bracketed}")
    with zipfile.ZipFile(REPORT) as archive:
        media = [n for n in archive.namelist() if n.startswith("word/media/")]
        require(len(media) >= 10, "Report has too few embedded figures")
        require("word/footer1.xml" in archive.namelist(), "Report footer/page field missing")
        archive.testzip()
    return {"paragraphs": len(doc.paragraphs), "tables": len(doc.tables), "figures": len(media)}


def validate_ppt() -> dict:
    require(PPT.exists() and PPT.stat().st_size > 500_000, "Presentation missing or unexpectedly small")
    with zipfile.ZipFile(PPT) as archive:
        bad = archive.testzip()
        require(bad is None, f"Corrupt PPT member: {bad}")
        slides = sorted(n for n in archive.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", n))
        media = [n for n in archive.namelist() if n.startswith("ppt/media/")]
        charts = [n for n in archive.namelist() if re.fullmatch(r"ppt/(?:slides/)?charts/chart\d+\.xml", n)]
        require(len(slides) == 16, f"Expected 16 slides, found {len(slides)}")
        require(len(media) >= 4, "Presentation missing embedded screenshots/charts")
        require(len(charts) >= 1, "Presentation missing editable evaluation chart")
    text = xml_text(PPT, r"ppt/slides/slide\d+\.xml")
    required = [
        "CampusEnergy AI", "Synthetic / simulated prototype", "4,320", "4,152", "624",
        "9.2585", "2.9061", "68.61%", "84.882", "12:00", "13:00", "EFFECTIVELY TIED",
        "No automatic", "WHAT-IF + EXPLAINABILITY", "44 / 44", "42.58 kWh error",
        "RESPONSIBLE AI + SDGS", "SDG 7", "IMPLEMENTED NOW", "PRODUCTION ROADMAP",
        "05  Error Analysis", "06  Responsible AI / About",
    ]
    for value in required:
        require(value.lower() in text.lower(), f"Presentation missing: {value}")
    lowered = text.lower()
    require("todo" not in lowered and "lorem ipsum" not in lowered, "Presentation contains unfinished text")
    require("[participant" not in lowered and "[team" not in lowered, "Presentation contains unresolved placeholders")
    return {"slides": len(slides), "media": len(media), "charts": len(charts)}


def validate_markdown() -> dict:
    files = [ROOT / "README.md", ROOT / "docs" / "Demo_Guide.md", ROOT / "docs" / "Viva_QA.md"]
    combined = "\n".join(p.read_text(encoding="utf-8") for p in files)
    for value in ["4,320", "4,152", "624", "9.2585", "2.9061", "68.61", "84.882", "12:00", "13:00", "effectively tied", "synthetic"]:
        require(value.lower() in combined.lower(), f"Markdown package missing canonical value: {value}")
    require("automatic equipment control" not in combined.lower() or "no automatic equipment control" in combined.lower(), "Control boundary ambiguous")
    return {"files": [str(p.relative_to(ROOT)) for p in files]}


def main() -> None:
    output = {
        "report": validate_report(),
        "presentation": validate_ppt(),
        "markdown": validate_markdown(),
        "canonical_source": {
            "selected_model": METRICS["selection"]["selected_model"],
            "test_rows": METRICS["final_test"]["test_rows"],
            "forecast_rows": METRICS["future_forecast"]["rows"],
        },
    }
    print(json.dumps(output, indent=2))


if __name__ == "__main__":
    main()
